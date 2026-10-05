import os
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from docx import Document
from docx.shared import Pt
HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HERE, "..", "inputs")
ss = getSampleStyleSheet(); pr = lambda t: Paragraph(t, ss["BodyText"])
def tbl(rows):
    t = Table(rows); t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8ECF1")), ("FONTSIZE", (0, 0), (-1, -1), 9)])); return t
doc = SimpleDocTemplate(os.path.join(OUT, "delegation_of_authority.pdf"), pagesize=A4, title="Delegation of authority matrix")
doc.build([Paragraph("Calloway Industrial Group: delegation of authority and procure-to-pay controls", ss["Title"]),
 Paragraph("Approval limits (per invoice, USD)", ss["Heading3"]),
 tbl([["Approver level", "Limit per invoice"], ["Approver L1", "10,000"], ["Approver L2", "50,000"], ["Approver L3", "250,000"]]),
 Spacer(1, 8),
 pr("An approver may not approve an invoice above their limit. Splitting a purchase into several invoices to stay under a limit is prohibited."),
 Paragraph("Purchase orders and receiving", ss["Heading3"]),
 pr("Invoices above USD 5,000 need a purchase order. Invoices of USD 5,000 or less may be paid without a purchase order with approver sign-off. "
    "Goods purchase orders (po_type GOODS) must have goods received and recorded in the goods-receipt file before payment. Services purchase orders (po_type SVC) are accepted on approver sign-off and have no goods receipt."),
 Paragraph("Vendor master and segregation of duties", ss["Heading3"]),
 pr("A vendor is identified by the vendor master. The ERP migration on 1 March 2025 re-keyed some vendors: a vendor with a legacy id and a current id is one vendor. "
    "The employee who creates a vendor must not enter that vendor's invoices. New vendors must pass TIN matching (tin_status matched)."),
 Paragraph("Bank account changes", ss["Heading3"]),
 pr("Every bank account change needs a call-back to a known vendor contact and a ticket reference (callback_verified Y). A change without call-back verification must not be followed by a payment of USD 25,000 or more to the new account within 14 days."),
 Paragraph("Blocked vendors", ss["Heading3"]),
 pr("No payment may be released to a vendor while it is blocked, unless the CFO approves an override and it is logged (block_overrides)."),
 Spacer(1, 10), Paragraph("Synthetic exercise document.", ss["Italic"])])

d = Document(); d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10.5)
d.add_heading("Forensic AP review programme (Internal Audit, v2)", 1)
d.add_paragraph("Scope: all accounts-payable activity from 1 January 2024 to 30 September 2025. Every case that meets a definition below is reported, with no minimum size. "
                "Exposure means money paid out that should not have been, after netting anything already recovered.")
d.add_heading("Tests and how exposure is measured", 2)
for t in [
 "1. Duplicate payment. The same vendor (legacy and current ids are one vendor), the same invoice number once non-digits and leading zeros are ignored, and the same amount, paid twice. Exposure is the later payment, unless that payment was voided or the money came back as a refund (payment_type REFUND_IN).",
 "2. Employee-linked vendor. A vendor whose bank account matches an employee's payroll account, or whose address and phone match an employee's home address and phone, and which is not on the related-party register. A similar name alone is not enough. Exposure is every cleared payment to the vendor.",
 "3. Split approvals. Two or more invoices to the same vendor, approved by the same approver within 5 calendar days, each between 90% and 100% of that approver's limit. Exposure is the total of the invoices in the cluster.",
 "4. Receipts shortfall. A goods purchase order where the goods received are less than 80% of the invoiced amount. Services purchase orders are excluded. Exposure is invoice amount minus goods received minus credit memos issued against that purchase order.",
 "5. Bank change diversion. A bank account change without call-back verification, followed within 14 days by a payment of USD 25,000 or more to the new account. Exposure is the payments sent to the new account.",
 "6. Shell vendor. A vendor created by an employee who also entered at least 80% of that vendor's invoices, whose TIN status is not matched. Exposure is every cleared payment to the vendor.",
 "7. Contract overrun. Cleared payments on invoices that reference a contract exceed 110% of the contract's cap, where the cap is the original value or the latest amendment. Exposure is total paid minus that cap.",
 "8. Blocked vendor. A cleared payment dated on or after the day a vendor was blocked and before any reactivation, with no logged override. Exposure is those payments.",
]: d.add_paragraph(t)
d.add_heading("What to do about each type", 2)
d.add_paragraph("Freeze (hold every payment to the vendor pending investigation) any vendor with confirmed exposure under test 2, 4, 5 or 6. Tests 1, 3, 7 and 8 are control failures to correct and recover; they do not by themselves justify freezing the vendor.")
d.add_paragraph("Report voided, refunded, credited, disclosed or overridden items as cleared rather than as exposure.")
d.save(os.path.join(OUT, "audit_programme_v2.docx"))
open(os.path.join(OUT, "change_log.txt"), "w").write("""# Calloway AP / ERP change log (informal)
2024-01-02 | AP | New procure-to-pay controls effective (see delegation of authority)
2024-06-14 | IT | Treasury payment export now writes dates as MM/DD/YYYY
2025-03-01 | ERP | Migration cut-over. About a third of vendors re-keyed from V##### to VN-######; invoices re-entered from the old system carry the vendor's own number in whatever format was typed
2025-03-01 | ERP | Treasury export unchanged; bank accounts are typed in several formats
2025-04-22 | AP | Finance asked AP to clean up duplicate vendor records after the migration
2025-07-30 | HR | HR extract refreshed; payroll accounts exported separately
""")
open(os.path.join(OUT, "data_notes.md"), "w").write("""# Data notes (Internal Audit, 1 Oct 2025)

Extracts cover 1 Jan 2024 - 30 Sep 2025. Synthetic data.

- `ap_invoices.csv` - AP invoice register (invoices and credit memos). `vendor_id` is the id in force on the invoice date. `invoice_no` is as typed.
- `payments.csv` - treasury payments. `invoice_ref` is the invoice number as typed on the payment run. Dates are MM/DD/YYYY. `payment_type` REFUND_IN is money received back.
- `vendor_master.xlsx` - vendors (current and legacy ids), status history, related-party register, block overrides.
- `bank_change_requests.json` - bank account change requests with call-back flags.
- `purchase_orders.csv`, `goods_receipts.csv` - procurement and receiving.
- `hr_employees.csv`, `hr_payroll_accounts.json` - HR and payroll bank accounts.
- `contracts_register.xlsx` - contracts and amendments.
- `delegation_of_authority.pdf`, `audit_programme_v2.docx`, `change_log.txt` - as named.
""")
print("docs ok")
