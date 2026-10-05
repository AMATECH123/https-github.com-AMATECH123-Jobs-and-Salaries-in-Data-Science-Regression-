import os, json
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from docx import Document
from docx.shared import Pt
HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HERE, "..", "inputs"); T = json.load(open(os.path.join(HERE, "world_truth.json")))
ss = getSampleStyleSheet(); pr = lambda t: Paragraph(t, ss["BodyText"])

# holdout design
doc = SimpleDocTemplate(os.path.join(OUT, "holdout_test_design.pdf"), pagesize=A4, title="Geo holdout test, Feb-Apr 2025")
doc.build([Paragraph("Geo holdout test: brand search and retargeting (Feb - Apr 2025)", ss["Title"]),
 pr("Why: both channels report a lot of conversions but we suspected many of those people would have contacted us anyway. We paused each one in a set of service-area clusters and left everything else running."),
 Paragraph("Design", ss["Heading3"]),
 pr("Test window: 3 Feb 2025 to 27 Apr 2025 (12 weeks). Reference period before the test: 11 Nov 2024 to 2 Feb 2025 (12 weeks)."),
 pr("Brand search ads were switched off in clusters: " + ", ".join(T["T_brand"]) + "."),
 pr("Retargeting was switched off in clusters: " + ", ".join(T["T_retarget"]) + "."),
 pr("Control: all other clusters, where both channels kept running as normal: " + ", ".join(T["control"]) + ". Every other channel kept running in every cluster."),
 pr("The pause clusters were picked by the regional manager from the busier markets so the effect would be easy to see. Cluster membership by zip code is in geo_map.csv."),
 pr("Spend for the two channels fell during the test because the paused clusters were not being served."),
 Spacer(1, 12), Paragraph("Synthetic exercise document.", ss["Italic"])])

# commission plan
doc = SimpleDocTemplate(os.path.join(OUT, "commission_and_cost_plan.pdf"), pagesize=A4, title="Commission and job cost plan")
t = Table([["Job revenue band", "Sales commission rate"], ["First $1,000", "5%"], ["$1,000 to $3,000", "8%"], ["Above $3,000", "12%"]]); t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8ECF1"))]))
doc.build([Paragraph("Commission and job cost plan (2025)", ss["Title"]),
 pr("Sales commission is paid on gross job revenue at completion, band by band (marginal rates, not a single rate for the whole job). It is not clawed back if the job is later refunded."),
 t, Spacer(1, 10),
 pr("Technician labour is costed at a burdened USD 48 per hour on the hours recorded in the job sheet."),
 pr("Card and financing processing costs 2.9% of the payment plus USD 0.30 per payment and is not returned on refunds."),
 pr("Materials are as recorded in the job sheet. Refunds and chargebacks are taken in full against the job's contribution."),
 Spacer(1, 12), Paragraph("Synthetic exercise document.", ss["Italic"])])

# measurement policy
d = Document(); st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
d.add_heading("Marketing measurement standard (Harbor & Pine, v2)", 1)
d.add_paragraph("Used whenever we decide where to put marketing money. It exists because three dashboards gave three different answers last year.")
d.add_heading("1. What counts as a lead", 2)
d.add_paragraph("A lead is a unique person. Two records are the same person if they share a normalised email, or share a phone number and a surname. "
                "Form submissions and contact records count; a call counts only if it was answered and lasted at least 30 seconds. "
                "A person's channel is the channel of the earliest touch recorded for them in any system.")
d.add_heading("2. What a lead is worth", 2)
d.add_paragraph("Contribution per won job = revenue minus materials, technician labour, sales commission, card fees, and refunds and chargebacks (see the cost plan). "
                "Within a channel, a lead is worth the same whether or not that person would have contacted us anyway.")
d.add_heading("3. Which period describes current performance", 2)
d.add_paragraph("Use leads whose first touch falls in the last 13 weeks (1 Jul - 30 Sep 2025). Conversion moved after the sales process changed in April, so older leads do not describe today's conversion.")
d.add_heading("4. Where the next dollar goes", 2)
d.add_paragraph("We rank channels by what the next marketing dollar returns, in contribution, at each channel's current weekly run-rate (average of the last 8 weeks). "
                "How volume responds to spend should be learned from the full 2024-2025 history.")
d.add_paragraph("Brand search and retargeting were put through a geo holdout test. The other channels were not tested and are treated as fully incremental.")
d.add_heading("5. Recommendation", 2)
d.add_paragraph("Name the one channel that should receive the next dollars, with the figures behind it.")
d.save(os.path.join(OUT, "measurement_policy_v2.docx"))

open(os.path.join(OUT, "change_log.txt"), "w").write("""# Harbor & Pine ops change log (informal)
2024-01-02 | GHL pipeline set up; web forms and call tracking numbers live
2024-03-10 | US daylight saving started; call tracking vendor reports in UTC
2024-06-28 | GHL export switched from local (Central) time to ISO UTC timestamps from 1 Jul 2024
2024-09-05 | Added a second Meta instant-form with a shorter questionnaire
2025-02-03 | Geo holdout test started (see design note)
2025-04-14 | New lead-qualification script launched for Meta and display leads; close rates moved after this date
2025-04-28 | Geo holdout test ended; all clusters back to normal
2025-06-10 | Rotated shared API keys
2025-08-01 | Front desk now logs phone leads in GHL the same day
""")
open(os.path.join(OUT, "data_notes.md"), "w").write("""# Data notes (marketing ops, 1 Oct 2025)

Pulled on 1 Oct 2025. Leads cover 1 Jan 2024 - 30 Sep 2025.

- `ghl_contacts.csv` - GHL contact export. Timestamps before 1 Jul 2024 are Central local time (MM/DD/YYYY HH:MM); later ones are ISO UTC. `source` is whatever the form or front desk typed.
- `call_tracking.json` - call tracking log. `started_at` is epoch milliseconds UTC. Each tracking number is tied to one channel in `channel_map.csv`.
- `airtable_deals.json` - Airtable deals. Each deal points at a contact by email, phone and/or GHL contact id, whichever the rep had.
- `jobs_completed.xlsx` - completed jobs, one sheet per quarter. Most rows carry the Airtable deal id; some do not.
- `payments.csv` - payments, refunds and chargebacks by job.
- `ad_spend_daily.csv` - platform spend by channel by day. `platform_reported_conversions` is what each platform's own dashboard claims.
- `season_index.csv` - weekly demand index from the market data service (1.0 = average).
- `geo_map.csv`, `holdout_test_design.pdf`, `commission_and_cost_plan.pdf`, `measurement_policy_v2.docx`, `change_log.txt` - as named.
- Figures in this folder are synthetic.
""")
print("docs ok")
