import os, pickle, json
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from docx import Document
from docx.shared import Pt
HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HERE, "..", "inputs")
P = pickle.load(open(os.path.join(HERE, "truth.pkl"), "rb")); prior = P["prior"]
ss = getSampleStyleSheet(); pr = lambda t: Paragraph(t, ss["BodyText"])
def tbl(rows):
    t = Table(rows); t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8ECF1")), ("FONTSIZE", (0, 0), (-1, -1), 9), ("ALIGN", (1, 1), (-1, -1), "RIGHT")])); return t
f = lambda x: f"{x:,.0f}"
doc = SimpleDocTemplate(os.path.join(OUT, "prior_year_workpapers.pdf"), pagesize=A4, title="Prior-year tax workpapers")
doc.build([Paragraph("Brightwater Precision Components, Inc. - prior-year tax workpaper summary (as filed)", ss["Title"]),
 Paragraph("1. Research and experimental expenditures capitalised under IRC section 174 (tax years beginning after 2021)", ss["Heading3"]),
 tbl([["Tax year", "Domestic research (USD)", "Foreign research (USD)"], ["2022", f(prior["re_dom_2022"]), f(prior["re_for_2022"])], ["2023", f(prior["re_dom_2023"]), f(prior["re_for_2023"])]]),
 pr("Amounts are the total costs paid or incurred in each year that were capitalised and are being amortised on the returns as filed. 2024 amounts are not shown here."),
 Spacer(1, 8), Paragraph("2. Net operating loss carryforward", ss["Heading3"]),
 pr(f"Federal NOL carryforward available at 1 January 2024: USD {f(P['NOL_BAL'])}. It arose in tax year 2020 (after 2017) and has not been used."),
 Paragraph("3. Section 163(j) disallowed business interest carryforward", ss["Heading3"]),
 pr("Disallowed business interest expense carried forward from 2023 to 2024: USD 450,000."),
 Paragraph("4. Gross receipts and entity facts", ss["Heading3"]),
 tbl([["Year", "Gross receipts (USD)"], ["2021", "98,400,000"], ["2022", "109,600,000"], ["2023", "118,900,000"]]),
 pr("C corporation, calendar tax year, accrual method. Majority shareholder: Ronald Hale (CEO), 60%; the remainder is held by a private equity fund. Financial statements are audited each year."),
 Spacer(1, 10), Paragraph("Synthetic exercise document.", ss["Italic"])])
d = Document(); d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10.5)
d.add_heading("CFO memo: instructions for the 2024 federal return", 1)
d.add_paragraph("Date: 12 September 2025. The 2024 return is on extension (due 15 October 2025). Please prepare the federal income tax computation and tell me what we owe.")
d.add_heading("Elections and policies", 2)
for t in [
 "Take every deduction and election the tax code allows for 2024. Do not elect out of bonus depreciation.",
 "Make the section 179 election on the following assets, in this order, up to the maximum the law allows: (1) CNC machining center (Haas VF-10), up to its full cost; (2) Fiber laser cutting system, up to its full cost. Do not elect section 179 on any other asset.",
 "Make the de minimis safe harbor election (Treas. Reg. 1.263(a)-1(f)) for the year. Our financial statements are audited. For books we capitalise purchases of USD 1,800 and above.",
 "For research costs, count only employee wages (based on the percentage of time each person spent on qualified research, in the payroll register), contractor invoices and prototype materials. We do not allocate payroll taxes, benefits or overhead to research.",
 "The 2023 and prior research and carryforward balances are in the prior-year workpapers. Estimated and extension payments are in tax_payments.csv.",
 "Ignore any estimated tax underpayment penalty and any state tax effects beyond what is already in the ledger. Do not compute research credits.",
]: d.add_paragraph(t, style="List Bullet")
d.save(os.path.join(OUT, "cfo_election_memo.docx"))
open(os.path.join(OUT, "change_log.txt"), "w").write("""# Brightwater finance systems change log
2024-01-02 | ERP | Chart of accounts unchanged for 2024
2024-07-01 | ERP | Upgrade: post_date now exports as MM/DD/YYYY; some account names shortened (codes unchanged)
2024-12-31 | FA | Fixed-asset subledger closed; 2024 additions entered with in-service dates
2025-02-14 | Payroll | Payroll register refreshed with R&D time percentages from engineering managers
2025-03-04 | Tax | Prior-year accrual schedule rolled forward
""")
open(os.path.join(OUT, "data_notes.md"), "w").write("""# Data notes (Tax, 12 Sep 2025)

Tax year 2024. Synthetic data.

- `gl_transactions.csv` - 2024 general ledger detail (debit/credit by line). Revenue is a daily summary; federal income tax expense (account 8100) is the book provision.
- `chart_of_accounts.xlsx` - accounts and types.
- `fixed_assets.xlsx` - all assets with book and tax data, plus 2024 disposals. For assets acquired before 2024 the prior workpapers already hold the 2024 tax depreciation. 2024 additions have no tax depreciation yet.
- `payroll_register.csv` - employees, annual wages, location and share of time on qualified research.
- `contractor_invoices.json` - 2024 contractor invoices booked to account 6530, with vendor country and description.
- `accrual_schedule.xlsx` - bonus and PTO accruals for 2024 and 2023 with payment dates.
- `tax_payments.csv`, `prior_year_workpapers.pdf`, `cfo_election_memo.docx`, `change_log.txt` - as named.
""")
print("docs ok")
