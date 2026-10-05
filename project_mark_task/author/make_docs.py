import os
from params import *
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from docx import Document
from docx.shared import Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "inputs")
ss = getSampleStyleSheet()

def tbl(data, widths=None):
    t = Table(data, colWidths=widths)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8ECF1")),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 9),
                           ("ALIGN", (1, 1), (-1, -1), "RIGHT")]))
    return t

# ---------------- pricing sheet
doc = SimpleDocTemplate(os.path.join(OUT, "pricing_quotes_sep2025.pdf"), pagesize=A4, title="Platform pricing quotes - Sep 2025")
s = [Paragraph("Automation platform pricing quotes (received Sep 2025)", ss["Title"]),
     Paragraph("Compiled by Ops from vendor emails and pricing pages. One plan is bought per platform for the whole renewal "
               "year (no mid-year plan changes). Billing and usage rules that apply to all three are in the contract summary.", ss["BodyText"]),
     Spacer(1, 10), Paragraph("Make (billed in EUR) - unit: operations per month", ss["Heading3"]),
     tbl([["Plan allowance (ops/mo)", "Monthly fee (EUR)", "Overage (EUR per 1,000 ops)"]] +
         [[f"{a:,}", f"{f:,}", f"{o:.2f}"] for a, f, o in MAKE_TIERS]),
     Spacer(1, 10), Paragraph("Zapier (billed in USD) - unit: tasks per month", ss["Heading3"]),
     tbl([["Plan allowance (tasks/mo)", "Monthly-billing list price (USD)"]] + [[f"{a:,}", f"{p:,}"] for a, p in ZAP_TIERS]),
     Paragraph("Prices above are the month-to-month list price. Annual billing terms are in the contract summary.", ss["BodyText"]),
     Spacer(1, 10), Paragraph("n8n Cloud (billed in EUR) - unit: workflow executions per month", ss["Heading3"]),
     tbl([["Plan allowance (executions/mo)", "Monthly fee (EUR), annual prepay"]] + [[f"{a:,}", f"{p:,}"] for a, p in N8N_TIERS]),
     Spacer(1, 14), Paragraph("Illustrative figures for an internal exercise; not actual vendor pricing.", ss["Italic"])]
doc.build(s)

# ---------------- contract summary
doc = SimpleDocTemplate(os.path.join(OUT, "contract_terms_summary.pdf"), pagesize=A4, title="Contract terms summary")
pr = lambda t: Paragraph(t, ss["BodyText"])
s = [Paragraph("Contract terms summary - renewals on 1 November 2025", ss["Title"]),
     Paragraph("Term and billing", ss["Heading3"]),
     pr("All three subscriptions renew on 1 November 2025 for a 12-month term (Nov 2025 - Oct 2026). Usage allowances reset on the 1st "
        "of each calendar month; unused allowance does not roll over."),
     Paragraph("Make", ss["Heading3"]),
     pr("Make has notified a list-price increase of 14% effective 1 January 2026. It applies to monthly plan fees and to overage rates "
        "for every plan, including plans bought at this renewal. Overage is billed monthly in blocks of 1,000 operations, rounded up."),
     Paragraph("Zapier", ss["Heading3"]),
     pr("Annual billing gives a 20% discount on the plan list price. The discount does not apply to overage. Overage is billed monthly "
        "per block of 1,000 tasks, rounded up, at 1.25 x the plan's list price per 1,000 allowance tasks "
        "(list price divided by allowance in thousands)."),
     Paragraph("n8n", ss["Heading3"]),
     pr("The quoted price is the annual-prepay price. There is no overage: executions above the monthly allowance are not run. "
        "The plan therefore has to cover the busiest month of the year."),
     Paragraph("Currency", ss["Heading3"]),
     pr("Make and n8n invoice in EUR. Zapier invoices in USD. Northgate budgets in EUR."),
     Spacer(1, 14), Paragraph("Illustrative summary for an internal exercise.", ss["Italic"])]
doc.build(s)

# ---------------- costing policy docx
d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
d.add_heading("Usage and Capacity Costing Standard (Finance, v3)", 1)
d.add_paragraph("Applies to any comparison of automation platform cost at Northgate. Written after the 2024 renewal, when three exports "
                "disagreed with each other and the decision was made on the wrong totals.")
d.add_heading("1. Unit of measure: the business run", 2)
d.add_paragraph("Platform invoices are not comparable because each vendor counts something different (operations, tasks, executions). "
                "Compare platforms on business runs. A business run is one first-attempt, successful, production execution of a workflow "
                "for one trigger event. Failures, retries, replays, manual or test runs and scratch workflows are not business runs. "
                "Internal Northgate workflows are real consumption and count. Every event is counted once, however many systems or "
                "exports happened to record it.")
d.add_heading("2. Billable units on a target platform", 2)
d.add_paragraph("Make bills operations (modules executed per run), Zapier bills tasks (billable steps per run), n8n bills executions "
                "(one per run). Per-run counts for the build that is currently deployed are in the workflow catalog.")
d.add_heading("3. Forward 12-month cost", 2)
for t in [
    "The forward year is Nov 2025 - Oct 2026 and repeats, month for month, the trailing window Oct 2024 - Sep 2025, for clients we still have.",
    "Subscription: one plan per platform for the year, chosen to minimise plan fees plus overage under the contract terms.",
    "Northgate budgets in EUR. Use the most recent published rate for the latest month in fx_rates.csv.",
]:
    d.add_paragraph(t, style="List Bullet")
d.add_heading("4. One-off migration effort", 2)
d.add_paragraph("To standardise on a platform, every in-scope workflow that is not already running on it must be rebuilt. "
                f"Labour is costed at EUR {RATE_EUR_PER_HOUR} per hour.")
t = d.add_table(rows=1, cols=3); t.style = "Light Grid Accent 1"
t.rows[0].cells[0].text = "Target platform"; t.rows[0].cells[1].text = "Base hours per workflow"; t.rows[0].cells[2].text = "Extra hours per node"
for p, lab in [("zapier", "Zapier"), ("make", "Make"), ("n8n", "n8n")]:
    r = t.add_row().cells; r[0].text = lab; r[1].text = f"{HOURS[p][0]:.1f}"; r[2].text = f"{HOURS[p][1]:.2f}"
d.add_paragraph(f"Nodes come from the workflow catalog. If a workflow needs a connector that the target platform does not support "
                f"(connector_support.csv), add {CUSTOM_BUILD_HOURS} hours of custom build per unsupported connector.")
d.add_heading("5. Decision rule", 2)
d.add_paragraph("Compare platforms on total 12-month cost in EUR = subscription (fees + overage) + one-off migration labour. "
                "Recommend the lowest and show the components.")
d.save(os.path.join(OUT, "costing_policy_v3.docx"))

# ---------------- data notes
open(os.path.join(OUT, "data_notes.md"), "w").write("""# Data notes (Ops, 2 Oct 2025)

Exports pulled on 2 Oct 2025 covering 1 Oct 2024 - 30 Sep 2025.

- `zapier_usage_daily.csv` - Zapier usage by zap by day, from three admin exports. Dates are UTC.
  `runs_success` = first-attempt successful runs. `runs_replayed` = replays of earlier errored runs. `test_runs` = runs from the editor.
  `tasks_billed` is what Zapier counted against our plan.
- `make_operations_daily.csv` - Make usage by scenario by day. Semicolon separated. Our Make org is set to German formatting.
  `Executions` is all executions that day; `Failed`, `Retries` and `Manual runs` are subsets of it. `Operations` is what Make billed.
- `n8n_executions.json` - execution log from the n8n API. `startedAt` is in the instance's local time.
- `id_crosswalk.csv` - native platform id to workflow_id. Blank workflow_id means a scratch item.
- `workflow_catalog.xlsx` - workflows (platform column last refreshed March 2025) and per-version module/step counts.
- `client_master.xlsx` - client list from the CRM. Internal workflows use client_id INTERNAL, which is not in this list.
- `connector_support.csv`, `fx_rates.csv`, `incident_log.csv`, `change_log.txt` - as named.
- Over the year a few zaps and scenarios were rebuilt or re-created by hand, so the crosswalk may not have every newest id.
- PDFs and the costing standard are as received / as published by Finance.
- Figures in this folder are synthetic.
""")
print("docs ok")
